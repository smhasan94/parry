#!/usr/bin/env python3
"""Render .agent/blockers.jsonl into shareable formats.

    render-blockers.py                    # all three formats into .agent/export/
    render-blockers.py --risk high        # only high-risk entries
    render-blockers.py --since 2026-08-01
    render-blockers.py --format md        # md | csv | xlsx | all

The JSONL is the record of truth. These outputs are disposable and regenerated.
"""

import argparse
import csv
import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path

COLUMNS = [
    ("date", "Date", 20),
    ("blocker", "Gap / Blocker", 60),
    ("resolution", "Resolution decided", 60),
    ("prev_commit", "Commit before fix", 16),
    ("risk", "Risk", 10),
    ("risk_explanation", "Risk explanation", 70),
    ("task", "Task", 22),
    ("worker", "Worker", 10),
    ("branch", "Branch", 22),
    ("id", "ID", 22),
]

RISK_ORDER = {"high": 0, "medium": 1, "low": 2}


def find_agent_dir(start: Path) -> Path:
    for d in [start, *start.parents]:
        if (d / ".agent").is_dir():
            return d / ".agent"
    sys.exit("render-blockers: no .agent directory found from %s" % start)


def load(path: Path, since=None, risk=None, task=None):
    if not path.exists():
        sys.exit(f"render-blockers: {path} does not exist yet")
    rows, bad = [], 0
    for n, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        line = line.strip()
        if not line:
            continue
        try:
            r = json.loads(line)
        except json.JSONDecodeError:
            bad += 1
            print(f"  warning: line {n} is not valid JSON, skipped", file=sys.stderr)
            continue
        if since and r.get("date", "") < since:
            continue
        if risk and r.get("risk", "").lower() != risk.lower():
            continue
        if task and task not in (r.get("task") or ""):
            continue
        rows.append(r)
    if bad:
        print(f"  {bad} malformed line(s) skipped", file=sys.stderr)
    return rows


def cell(r, key):
    v = r.get(key, "")
    return "" if v is None else str(v)


def write_md(rows, out: Path):
    lines = [
        "# Blocker and resolution log",
        "",
        f"Generated {datetime.now(timezone.utc).astimezone().isoformat(timespec='seconds')} "
        f"from `.agent/blockers.jsonl` — {len(rows)} entr{'y' if len(rows)==1 else 'ies'}.",
        "",
        "Risk is the confidence in the **resolution**, not the severity of the blocker. "
        "See `.agent/docs/blocker-schema.md`.",
        "",
    ]
    counts = {k: sum(1 for r in rows if r.get("risk") == k) for k in ("high", "medium", "low")}
    lines += [
        f"**High {counts['high']} · Medium {counts['medium']} · Low {counts['low']}**",
        "",
        "---",
        "",
    ]
    for r in rows:
        badge = {"high": "🔴 HIGH", "medium": "🟡 MEDIUM", "low": "🟢 LOW"}.get(
            r.get("risk", ""), r.get("risk", "?")
        )
        date = cell(r, "date").replace("T", " ")[:19]
        lines += [
            f"## {date} — {badge}",
            "",
            f"**Blocker.** {cell(r, 'blocker')}",
            "",
            f"**Resolution.** {cell(r, 'resolution')}",
            "",
            f"**Risk ({cell(r, 'risk')}).** {cell(r, 'risk_explanation')}",
            "",
            f"`commit before fix: {cell(r, 'prev_commit')}`"
            + (f" · task `{cell(r, 'task')}`" if cell(r, "task") else "")
            + (f" · branch `{cell(r, 'branch')}`" if cell(r, "branch") else "")
            + (f" · worker `{cell(r, 'worker')}`" if cell(r, "worker") else ""),
            "",
            "---",
            "",
        ]
    out.write_text("\n".join(lines), encoding="utf-8")
    return out


def write_csv(rows, out: Path):
    with out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow([h for _, h, _ in COLUMNS])
        for r in rows:
            w.writerow([cell(r, k) for k, _, _ in COLUMNS])
    return out


def write_xlsx(rows, out: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Blocker log"

    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3864")
    body_font = Font(name="Arial", size=10)
    mono_font = Font(name="Arial", size=10, color="555555")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    risk_fill = {
        "high": PatternFill("solid", fgColor="F8CBAD"),
        "medium": PatternFill("solid", fgColor="FFE699"),
        "low": PatternFill("solid", fgColor="C6E0B4"),
    }

    for c, (_, head, width) in enumerate(COLUMNS, 1):
        cl = ws.cell(row=1, column=c, value=head)
        cl.font, cl.fill, cl.border = header_font, header_fill, border
        cl.alignment = Alignment(vertical="center", horizontal="left")
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 22

    for i, r in enumerate(rows, start=2):
        for c, (key, _, _) in enumerate(COLUMNS, 1):
            val = cell(r, key)
            if key == "date":
                val = val.replace("T", " ")[:19]
            cl = ws.cell(row=i, column=c, value=val)
            cl.font = mono_font if key in ("prev_commit", "id") else body_font
            cl.border = border
            cl.alignment = Alignment(
                vertical="top",
                wrap_text=key in ("blocker", "resolution", "risk_explanation"),
            )
            if key == "risk":
                cl.fill = risk_fill.get(val.lower(), PatternFill())
                cl.font = Font(name="Arial", size=10, bold=True)
                cl.alignment = Alignment(vertical="top", horizontal="center")
        ws.row_dimensions[i].height = 46

    last = max(len(rows) + 1, 2)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"

    # Summary sheet. Formulas rather than baked numbers, so the counts stay true
    # if someone filters or edits the log sheet.
    s = wb.create_sheet("Summary")
    s["A1"] = "Blocker log summary"
    s["A1"].font = Font(name="Arial", size=13, bold=True)
    s["A3"] = "Generated"
    s["B3"] = datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")
    s["A4"] = "Source"
    s["B4"] = ".agent/blockers.jsonl"
    s["A5"] = "Entries"
    s["B5"] = f"=COUNTA('Blocker log'!A2:A{last})"

    s["A7"] = "Risk"
    s["B7"] = "Count"
    for c in ("A7", "B7"):
        s[c].font = header_font
        s[c].fill = header_fill
    for n, level in enumerate(("high", "medium", "low"), start=8):
        s[f"A{n}"] = level
        s[f"A{n}"].fill = risk_fill[level]
        s[f"A{n}"].font = Font(name="Arial", size=10, bold=True)
        s[f"B{n}"] = f"=COUNTIF('Blocker log'!E2:E{last},\"{level}\")"
    s["A11"] = "Total"
    s["A11"].font = Font(name="Arial", size=10, bold=True)
    s["B11"] = "=SUM(B8:B10)"

    s["A13"] = (
        "Risk is confidence in the RESOLUTION, not severity of the blocker. "
        "High means the resolution rests on an assumption that could not be "
        "checked from inside the session, and is a request for review."
    )
    s["A13"].alignment = Alignment(wrap_text=True, vertical="top")
    s.merge_cells("A13:D17")
    s.column_dimensions["A"].width = 24
    s.column_dimensions["B"].width = 34
    for row in s.iter_rows(min_row=1, max_row=17, max_col=4):
        for c in row:
            if c.font.name is None or c.font.name != "Arial":
                c.font = Font(name="Arial", size=10, bold=c.font.bold)

    wb.save(out)
    return out


def main():
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--format", default="all", choices=["md", "csv", "xlsx", "all"])
    p.add_argument("--since", help="ISO date, e.g. 2026-08-01")
    p.add_argument("--risk", choices=["low", "medium", "high"])
    p.add_argument("--task", help="substring match on task id")
    p.add_argument("--out", help="output directory (default .agent/export)")
    p.add_argument("--log", help="path to blockers.jsonl")
    p.add_argument("--sort", default="date", choices=["date", "risk"])
    a = p.parse_args()

    agent = Path(os.environ["AGENT_DIR"]) if os.environ.get("AGENT_DIR") \
        else find_agent_dir(Path.cwd())
    log = Path(a.log) if a.log else agent / "blockers.jsonl"
    outdir = Path(a.out) if a.out else agent / "export"
    outdir.mkdir(parents=True, exist_ok=True)

    rows = load(log, since=a.since, risk=a.risk, task=a.task)
    if a.sort == "risk":
        rows.sort(key=lambda r: (RISK_ORDER.get(r.get("risk", ""), 9), r.get("date", "")))
    else:
        rows.sort(key=lambda r: r.get("date", ""))

    if not rows:
        print("render-blockers: no entries matched", file=sys.stderr)
        return 1

    made = []
    if a.format in ("md", "all"):
        made.append(write_md(rows, outdir / "blockers.md"))
    if a.format in ("csv", "all"):
        made.append(write_csv(rows, outdir / "blockers.csv"))
    if a.format in ("xlsx", "all"):
        try:
            made.append(write_xlsx(rows, outdir / "blockers.xlsx"))
        except ImportError:
            print("render-blockers: openpyxl not installed, skipping xlsx "
                  "(pip install openpyxl)", file=sys.stderr)

    print(f"{len(rows)} entries ->")
    for m in made:
        print(f"  {m}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
